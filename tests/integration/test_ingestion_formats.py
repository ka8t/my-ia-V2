"""Tests régression — formats d'ingestion (P0.2).

Couvre le bug observé : ``partition_xlsx() not available because dependencies
are not installed`` après upload d'un xlsx en prod. Les extras
``unstructured[xlsx,docx,pptx]`` doivent être présents dans l'image V2.

Ces tests valident :
  1. Les imports ``partition_{xlsx,docx,pptx}`` ne lèvent pas ``ImportError``.
  2. ``unstructured.partition.auto.partition`` route correctement les formats
     Office (sans tomber sur le ``MissingDependencyError``).
"""
from __future__ import annotations

import io

import pytest


def test_partition_xlsx_importable() -> None:
    """Le bug observé : partition_xlsx levait MissingDependencyError. Doit
    désormais s'importer sans erreur (extras [xlsx] installés)."""
    from unstructured.partition.xlsx import partition_xlsx  # noqa: F401


def test_partition_docx_importable() -> None:
    from unstructured.partition.docx import partition_docx  # noqa: F401


def test_partition_pptx_importable() -> None:
    from unstructured.partition.pptx import partition_pptx  # noqa: F401


def test_xlsx_end_to_end_parse() -> None:
    """End-to-end : génère un xlsx en mémoire et appelle partition() (auto-
    detect via mime type). Vérifie qu'on récupère au moins une cellule."""
    import openpyxl
    from unstructured.partition.auto import partition

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Header"
    ws["B1"] = "Value"
    ws["A2"] = "Foo"
    ws["B2"] = 42
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    elements = partition(
        file=buf,
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )
    assert len(elements) >= 1
    text = "\n".join(str(e) for e in elements)
    assert "Header" in text
    assert "Foo" in text
